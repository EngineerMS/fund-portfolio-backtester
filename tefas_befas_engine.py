import pandas as pd
from tefas import Crawler
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
import time

class FundMatrixGenerator:
    """
    A professional class that fetches TEFAS and BEFAS (Pension) fund prices via API, 
    filters them based on desired frequency and specific days, 
    and generates an Excel pivot matrix for investment simulations.
    """
    def __init__(self, start_date, end_date):
        self.tefas_api = Crawler()
        self.start_date = start_date
        self.end_date = end_date
        
    def _fetch_fund_data(self, fund_code, fund_type):
        """Fetches raw historical price data for a single fund within the given date range."""
        print(f"Fetching data for [{fund_code}]... Type: ({fund_type})")
        try:
            # Gentle delay to prevent WAF (Web Application Firewall) blocks
            time.sleep(0.4) 
            data = self.tefas_api.fetch(start=self.start_date, end=self.end_date, 
                                        name=fund_code, columns=["date", "code", "price"], kind=fund_type)
            df = pd.DataFrame(data)
            if not df.empty:
                df['date'] = pd.to_datetime(df['date'])
                return df
        except Exception as e:
            print(f"[ERROR] Failed to fetch data for {fund_code}: {e}")
        return pd.DataFrame()

    def generate_matrix(self, fund_list, fund_type, frequency='Monthly', target_day=1):
        """
        Versatile matrix generation engine.
        frequency: 'Monthly' or 'Daily'
        target_day: If 'Monthly' is selected, which day of the month should be targeted? (e.g., 1, 15)
        Note: If the target day is a holiday/weekend, the FIRST AVAILABLE TRADING DAY after it is automatically selected.
        """
        all_dataframes = []
        for fund in fund_list:
            df = self._fetch_fund_data(fund, fund_type)
            if not df.empty:
                all_dataframes.append(df)
                
        if not all_dataframes:
            return pd.DataFrame()
            
        df_combined = pd.concat(all_dataframes, ignore_index=True)
        df_combined = df_combined.sort_values(by=['code', 'date'])
        
        if frequency.lower() == 'monthly':
            # Group data by Year-Month
            df_combined['year_month'] = df_combined['date'].dt.to_period('M')
            
            # Option 1: Logic to find the specific day (or the first available trading day after it)
            # Filter days that are greater than or equal to the target day
            mask = df_combined['date'].dt.day >= target_day
            valid_days = df_combined[mask]
            
            # From the filtered list, pick the first encountered day for each month
            filtered_data = valid_days.groupby(['code', 'year_month']).first().reset_index()
            
            # Format the date for visual presentation (e.g., 01.01.2023)
            filtered_data['formatted_date'] = filtered_data['date'].dt.strftime("%d.%m.%Y")
            
        elif frequency.lower() == 'daily':
            # Option 2: Day-by-day raw data matrix
            filtered_data = df_combined.copy()
            filtered_data['formatted_date'] = filtered_data['date'].dt.strftime("%d.%m.%Y")
        else:
            raise ValueError("Frequency parameter must be either 'Monthly' or 'Daily'.")

        # Create Pivot Table (Rows: Funds, Columns: Dates)
        pivot_df = filtered_data.pivot(index='code', columns='formatted_date', values='price')
        
        # Sort columns (dates) chronologically
        column_order = pivot_df.columns.tolist()
        column_order.sort(key=lambda x: pd.to_datetime(x, format="%d.%m.%Y"))
        return pivot_df[column_order]

def export_to_excel(df_befas, df_tefas, befas_funds, tefas_funds, filename):
    """Exports Pandas Pivot tables to an Excel file with professional formatting."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Dynamic_Historical_Matrix"

    bold_font = Font(bold=True)
    center_align = Alignment(horizontal="center")

    all_dates = set()
    if not df_befas.empty: all_dates.update(df_befas.columns)
    if not df_tefas.empty: all_dates.update(df_tefas.columns)

    sorted_dates = list(all_dates)
    sorted_dates.sort(key=lambda x: pd.to_datetime(x, format="%d.%m.%Y"))

    row_num = 1
    
    # --- WRITE BEFAS SECTION ---
    ws.cell(row=row_num, column=1, value="BEFAS (PENSION)").font = bold_font
    for c_idx, date_str in enumerate(sorted_dates, start=2):
        c = ws.cell(row=row_num, column=c_idx, value=date_str)
        c.font = bold_font
        c.alignment = center_align
    row_num += 1

    for fund in befas_funds:
        ws.cell(row=row_num, column=1, value=fund).font = bold_font
        if not df_befas.empty and fund in df_befas.index:
            for c_idx, date_str in enumerate(sorted_dates, start=2):
                val = df_befas.at[fund, date_str] if date_str in df_befas.columns else None
                if pd.notna(val):
                    ws.cell(row=row_num, column=c_idx, value=val).number_format = '#,##0.0000'
        row_num += 1

    row_num += 2 

    # --- WRITE TEFAS SECTION ---
    ws.cell(row=row_num, column=1, value="TEFAS (MUTUAL)").font = bold_font
    for c_idx, date_str in enumerate(sorted_dates, start=2):
        c = ws.cell(row=row_num, column=c_idx, value=date_str)
        c.font = bold_font
        c.alignment = center_align
    row_num += 1

    for fund in tefas_funds:
        ws.cell(row=row_num, column=1, value=fund).font = bold_font
        if not df_tefas.empty and fund in df_tefas.index:
            for c_idx, date_str in enumerate(sorted_dates, start=2):
                val = df_tefas.at[fund, date_str] if date_str in df_tefas.columns else None
                if pd.notna(val):
                    ws.cell(row=row_num, column=c_idx, value=val).number_format = '#,##0.0000'
        row_num += 1

    ws.column_dimensions['A'].width = 16
    for i in range(2, len(sorted_dates) + 2):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 13

    wb.save(filename)
    print(f"\n[SUCCESS] Dynamic Excel file created: '{filename}'")

# ==========================================
# GITHUB USAGE EXAMPLE (USER AREA)
# ==========================================
if __name__ == "__main__":
    
    # 1. Set Environment and Dates
    engine = FundMatrixGenerator(start_date="2023-01-01", end_date="2023-12-31")
    
    befas_list = ['HES', 'BHT', 'AMZ']
    tefas_list = ['IIH', 'HVS', 'YAY']
    
    print("--- STARTING SCENARIO ENGINE ---\n")
    
    # Example 1: Fetch prices for the 15th of each month (or the next available trading day)
    print(">>> Running Monthly Engine (Target: 15th of each month)...")
    df_befas_monthly = engine.generate_matrix(befas_list, "EMK", frequency='Monthly', target_day=15)
    df_tefas_monthly = engine.generate_matrix(tefas_list, "YAT", frequency='Monthly', target_day=15)
    
    # Example 2: Fetch DAY-BY-DAY full capacity data for 2023 (Daily Reporting)
    # df_befas_daily = engine.generate_matrix(befas_list, "EMK", frequency='Daily')
    # df_tefas_daily = engine.generate_matrix(tefas_list, "YAT", frequency='Daily')
    
    # Export to Excel
    export_to_excel(df_befas_monthly, df_tefas_monthly, befas_list, tefas_list, "Monthly_15th_Matrix.xlsx")
